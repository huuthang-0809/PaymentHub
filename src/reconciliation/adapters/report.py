"""Exports of persisted result evidence, without raw payloads."""
import csv
import io
import json
from dataclasses import fields
from enum import Enum

from reconciliation.application.query import result_dict
from reconciliation.domain import ReconResult

COLUMNS = [field.name for field in fields(ReconResult)]
MONEY = {'system_amount_minor', 'provider_amount_minor', 'refunded_amount_minor', 'delta_amount'}


def export_report(results, format, data_as_of):
    rows = [{key: value.value if isinstance(value, Enum) else value
             for key, value in result_dict(row).items()} for row in results]
    if format == 'csv':
        stream = io.StringIO(newline='')
        writer = csv.DictWriter(stream, fieldnames=COLUMNS)
        writer.writeheader()
        for row in rows:
            # Treat user-controlled strings as text when opened in spreadsheets.
            writer.writerow({key: "'" + value if isinstance(value, str)
                             and value.lstrip().startswith(('=', '+', '-', '@')) else value
                             for key, value in row.items()})
        return stream.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8'
    if format == 'xlsx':
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Results'
        sheet.append(COLUMNS)
        for row in rows:
            sheet.append([str(row[key]) if key in MONEY and row[key] is not None
                          else row[key] for key in COLUMNS])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        metadata = workbook.create_sheet('Metadata')
        metadata.append(['data_as_of', str(data_as_of)])
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    import pyarrow as pa
    import pyarrow.parquet as pq
    from decimal import Decimal
    schema = pa.schema([
        pa.field(key, pa.decimal128(38, 0) if key in MONEY else
                 pa.int64() if key == 'source_row_number' else pa.string())
        for key in COLUMNS
    ])
    schema = schema.with_metadata({
        b'data_as_of': str(data_as_of).encode(),
        b'data_dictionary': json.dumps({
            field.name: {'type': str(field.type), 'nullable': True,
                         'description': ('Integer minor units; currency is stored per side.'
                                         if field.name in MONEY else
                                         'Persisted reconciliation evidence: ' + field.name)}
            for field in schema
        }).encode(),
    })
    for row in rows:
        for key in MONEY:
            if row[key] is not None:
                row[key] = Decimal(row[key])
    stream = io.BytesIO()
    pq.write_table(pa.Table.from_pylist(rows, schema=schema), stream)
    return stream.getvalue(), 'application/vnd.apache.parquet'
